"""Gerador da fixture mínima do Hias (Task 3, spike).

Cria um Hias sintético pequeno (BD2 com tabela Tabela1 + BD1 com tabela BD_1)
para exercitar o pipeline UNO rapidamente, antes/depois do arquivo real.
Espelha o helper _fazer_hias_minimo dos testes da Task 4. Usa openpyxl apenas
nesta fixture sintética — o arquivo REAL nunca passa pelo openpyxl.
"""
import sys

import openpyxl
from openpyxl.worksheet.table import Table

SAIDA = "/home/apppf/spike/entrada_minima.xlsx"


def gerar(caminho=SAIDA, n_linhas_bd2=12, com_bd1=True):
    """Gera o Hias mínimo: BD2 com N linhas + Tabela1; BD1 com BD_1 (opcional)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bd2 = wb.create_sheet("BD2")
    bd2.append(["Convênio", "Data", "A", "B", "C", "D", "E", "F", "G"])
    for i in range(1, n_linhas_bd2):
        bd2.append([f"CONV {i}", "01/01/2026", 1, 2, 3, 4, 5, 6, 7])
    bd2.add_table(Table(displayName="Tabela1", ref=f"A1:I{n_linhas_bd2}"))
    if com_bd1:
        bd1 = wb.create_sheet("BD1")
        for col in "ABCDEFGHIJKLMNOPQRSTUV":
            bd1[f"{col}1"] = f"hdr{col}"
        bd1.append(["117129", 1] + [""] * 20)
        bd1.add_table(Table(displayName="BD_1", ref="A1:V2"))
    wb.save(caminho)
    print("FIXTURE OK ->", caminho)


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else SAIDA
    gerar(caminho)
